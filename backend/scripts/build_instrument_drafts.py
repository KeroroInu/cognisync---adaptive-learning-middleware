"""Read reviewed source DOCX files and emit a reproducible, unpublished catalog to stdout.

Run with the bundled document Python (python-docx). Never writes the source documents
or database. Save its JSON output through the normal reviewed file-edit workflow.
"""

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document


def question(qid, text, kind="text", labels=(), **extra):
    return dict(
        id=qid,
        text=text,
        type=kind,
        required=False,
        options=[dict(value=str(i + 1), label=label) for i, label in enumerate(labels)],
        **extra,
    )


def condition(qid, *values):
    return dict(question_id=qid, values=list(values))


def paragraphs(file):
    return [p.text for p in Document(file).paragraphs if p.text.strip()]


def blocks(lines):
    result = {}
    current = None
    for line in lines:
        match = re.match(r"^(\d+)\.\s*(.*)", line)
        if match:
            current = int(match[1])
            result[current] = [match[2]]
        elif current and not re.match(r"^[一二三四五六] |^问卷到此|^接下来请所有", line):
            result[current].append(line)
    return result


def background(file):
    source = paragraphs(file)
    grouped = blocks(source)
    result = []
    programming = condition("q3", "1", "2", "3", "4")
    ai = condition("q8", "2", "3", "4", "5", "6", "7")
    ai_coding = condition("q11", "1")
    for n in range(1, 20):
        qid = f"q{n}"
        title, *body = grouped[n]
        conditions = (
            [programming]
            if 4 <= n <= 7
            else [ai]
            if 9 <= n <= 11
            else [ai_coding]
            if n in (12, 13)
            else []
        )
        if n == 6:
            q = question(qid, title, "choice", ["没有", "有", "记不清楚"], show_if=conditions)
            result.append(q)
            result.append(
                question(
                    "q6_methods",
                    "第6题补充：准备怎样完成任务？（可勾选多种方式）",
                    "multiple",
                    ["口头说明", "文字步骤", "图示", "其他"],
                    show_if=[condition(qid, "2")],
                )
            )
            result.append(
                question("q6_other", "第6题补充：其他方式", show_if=[condition("q6_methods", "4")])
            )
            continue
        if n == 10:
            result.append(question(qid, title, show_if=conditions, skip_label="不记得名称"))
            continue
        if n == 19:
            result.append(question(qid, title))
            continue
        labels = []
        for line in body:
            if "□" not in line:
                continue
            for raw in line.split("□")[1:]:
                label = re.sub(r"_+", "", raw).strip()
                label = re.sub(r"（请跳至第\d+题）", "", label)
                label = label.rstrip("：")
                if n == 4 and label.startswith("有："):
                    label = "有"
                labels.append(label)
        kind = "multiple" if n in (1, 2, 9, 12, 13, 14, 16, 17, 18) else "choice"
        q = question(qid, title, kind, labels, show_if=conditions)
        if n in (16, 17):
            q["max_selections"] = 3
        if n == 1:
            q["hint"] = source[1] + " AI指可以通过对话提供回答或帮助完成任务的人工智能工具。"
        if kind == "multiple":
            for o in q["options"]:
                if o["label"].startswith(("没有", "不确定", "尚未", "记不清", "暂未", "暂时没有")):
                    o["exclusive"] = True
        result.append(q)
        if n == 4:
            result.append(
                question("q4_detail", "第4题补充：作品名称或内容", show_if=[condition(qid, "2")])
            )
            result.append(
                question("q4_role", "第4题补充：你实际负责了什么？", show_if=[condition(qid, "2")])
            )
        for option in q["options"]:
            if option["label"] == "其他":
                result.append(
                    question(
                        f"{qid}_other",
                        f"第{n}题补充：其他",
                        show_if=[condition(qid, option["value"])],
                    )
                )
        if n == 12:
            result.append(
                question("q12_task", "第12题补充：当时的任务是什么？", show_if=[ai_coding])
            )
    for q in result:
        if "_" in q["id"]:
            q["parent_id"] = q["id"].split("_")[0]
    assert len([q for q in result if re.fullmatch(r"q\d+", q["id"])]) == 19
    return result


def scale(file, prefix, size):
    doc = Document(file)
    result = []
    for table in doc.tables:
        labels = [c.text.replace("\n", "") for c in table.rows[0].cells[2:]]
        for row in table.rows[1:]:
            values = [c.text for c in row.cells]
            result.append(question(f"{prefix}{int(values[0]):02d}", values[1], "choice", labels))
    assert len(result) == size
    result[0]["hint"] = next(p.text for p in doc.paragraphs if p.text.startswith("请"))
    return result


def exams(file):
    forms = {}
    form = None
    current = None
    for line in paragraphs(file):
        if "测试" in line and line.endswith(("A卷", "B卷")):
            form = line[-2]
            forms[form] = {}
            current = None
            continue
        match = re.match(r"^(\d+)\.\s*(.*)", line)
        if match:
            current = int(match[1])
            forms[form][current] = [match[2]]
        elif current and not re.match(r"^[一二三四]、|^[AB]卷|^班级：|^满分|^请独立", line):
            forms[form][current].append(line)
    output = {}
    for form, grouped in forms.items():
        result = []
        for n, parts in grouped.items():
            q = question(f"{form.lower()}{n:02d}", parts[0], "choice" if n <= 6 else "text")
            code = []
            extra = []
            for part in parts[1:]:
                if "A." in part and n <= 6:
                    for letter, label in re.findall(
                        r"([A-D])\.\s*(.*?)(?=\s+[A-D]\.|$)", part, re.S
                    ):
                        q["options"].append(dict(value=letter, label=f"{letter}. {label.strip()}"))
                elif "print(" in part:
                    code.append(part)
                else:
                    extra.append(part)
            q["code"] = "\n".join(code)
            q["hint"] = "\n".join(extra)
            if n == 1:
                q["hint"] = (
                    "请独立作答，不使用 AI、教材、搜索或程序运行工具。所有代码均为 Python 3；选择题只有一个正确选项。不会的题可以留空。"
                )
            assert n > 6 or len(q["options"]) == 4
            result.append(q)
        assert len(result) == 10
        output[form] = result
    return output


def build(root):
    paths = dict(
        background=root / "01_现状调查问卷/高中生编程学习与AI使用现状调查问卷_待审查稿.docx",
        ct=root / "05_计算思维测评/计算思维量表_20题.docx",
        cps=root / "06_CPS与编程问题解决评价/复杂编程问题解决自评问卷_审题稿.docx",
        exam=root / "03_编程知识测试/编程知识与程序理解测试_AB卷_学生用卷.docx",
    )
    data = dict(
        background=background(paths["background"]),
        ct=scale(paths["ct"], "ct", 20),
        cps=scale(paths["cps"], "cps", 16),
        **exams(paths["exam"]),
    )
    specs = [
        ("background", "pre", "现状调查 · 19个主问题", "background", 19),
        ("ct_pre", "pre", "CT20 · 前测", "ct", 20),
        ("ct_post", "post", "CT20 · 后测", "ct", 20),
        ("cps_pre", "pre", "CPS16 · 前测审题稿", "cps", 16),
        ("cps_post", "post", "CPS16 · 后测审题稿", "cps", 16),
        ("knowledge_a", "pre", "编程知识与程序理解 · A卷", "A", 10),
        ("knowledge_b", "post", "编程知识与程序理解 · B卷", "B", 10),
    ]
    notes = {
        "background": "自编背景调查，不计能力总分。19个主问题，其他说明及分支为子字段。保持Word跳题：无编程经历跳至8；从未用AI跳至14；未用AI编程跳至14。Q10勾选不记得名称保存为__skip__，不等于留空。",
        "ct": "Korkmaz与Bai（2019）20题中文版，非21题版。1—5对应从不至总是。16—20题拟按6−原分处理，来自题意核对，尚无作者计分程序确认。保留原分，缺答不记0；当前不自动评分。",
        "cps": "研究者自编四领域16题审题建议，尚未经用户/导师定题、专家审查或预试。不是Promma8，不是已验证量表，不用于正式施测；目前只供审阅。不自动评分。",
        "exam": "自编10题30分，A/B尚未证明等值，需审查与试做计时。独立作答，不使用AI、教材、搜索或运行工具；平台不能阻止本机其他工具，需教师组织。Q7—10按题保留多空文本，按独立评分标准人工核分。",
    }
    items = []
    for key, phase, name, content, count in specs:
        source_key = "exam" if content in ("A", "B") else content
        file = paths[source_key]
        digest = hashlib.sha256(file.read_bytes()).hexdigest()
        items.append(
            dict(
                key=key,
                main_questions=count,
                source_file=file.name,
                source_sha256=digest,
                draft=dict(
                    phase=phase,
                    name=f"待审阅｜{name}｜20261012",
                    class_name="预试班",
                    source=f"研究阶段：未分配（不是正式前测，也不是已施测预试）；题本代码：{key}；版本：20261012-draft1。来源：{file.name}；SHA256：{digest}。{notes[source_key]} 正式开放前确认样本、用途及版本；前后测分别开放，不混总分。",
                    questions=data[content],
                ),
            )
        )
    return dict(
        schema="classroom-instrument-catalog/1", version="20261012-draft1", instruments=items
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source_directory", type=Path)
    args = parser.parse_args()
    print(json.dumps(build(args.source_directory), ensure_ascii=False, indent=2))
